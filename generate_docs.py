import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.enum.text import PP_ALIGN

# --- Paths ---
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLOTS_DIR = os.path.join(BASE_DIR, "plots")
DOCX_PATH = os.path.join(BASE_DIR, "Project_Report.docx")
PPTX_PATH = os.path.join(BASE_DIR, "Project_Presentation.pptx")

# ==========================================
# Generate DOCX Report
# ==========================================
def create_docx_report():
    doc = Document()
    
    # Title
    title = doc.add_heading('Machine Learning Case Study: Chronic Kidney Disease Classification', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. Problem Statement & Objectives', level=1)
    doc.add_paragraph(
        "A healthcare organization wants to develop a Machine Learning system that classifies patients "
        "based on available clinical measurements and identifies whether they belong to a kidney-disease-related "
        "class (Chronic Kidney Disease or Not CKD).\n\n"
        "Objectives:\n"
        "1. Understand and clean the clinical dataset.\n"
        "2. Handle missing values appropriately.\n"
        "3. Perform Exploratory Data Analysis (EDA).\n"
        "4. Identify important predictors (Feature Engineering).\n"
        "5. Develop and compare ML models (Logistic Regression, KNN, Decision Tree, Random Forest, Gradient Boosting).\n"
        "6. Deploy a web application for real-time prediction."
    )

    doc.add_heading('2. Dataset Overview & Preprocessing', level=1)
    doc.add_paragraph(
        "The UCI Chronic Kidney Disease Dataset contains 400 patient records with 24 clinical features and 1 target variable. "
        "The dataset exhibited significant missing values (up to 38% in some features like Red Blood Cell count). "
        "To handle this, a two-pass imputation strategy was implemented: "
        "K-Nearest Neighbors (KNN) Imputation for numerical features and Mode Imputation for categorical features. "
        "This allowed us to retain all 400 patient records without data loss."
    )

    doc.add_heading('3. Exploratory Data Analysis (EDA)', level=1)
    doc.add_paragraph("Several visualizations were generated to understand feature distributions and class separation.")
    
    # Add EDA images if they exist
    eda_images = ["01_target_distribution.png", "04_numerical_distributions.png", "05_correlation_heatmap.png"]
    for img in eda_images:
        img_path = os.path.join(PLOTS_DIR, img)
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(5.0))

    doc.add_heading('4. Feature Importance', level=1)
    doc.add_paragraph(
        "Using Random Forest Feature Importances and Mutual Information scores, the most critical indicators for CKD were identified as:\n"
        "1. Hemoglobin (hemo)\n"
        "2. Packed Cell Volume (pcv)\n"
        "3. Specific Gravity (sg)\n"
        "4. Serum Creatinine (sc)\n"
        "5. Albumin (al)"
    )
    
    img_path = os.path.join(PLOTS_DIR, "08_feature_importance.png")
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(5.5))

    doc.add_heading('5. Model Development & Comparison', level=1)
    doc.add_paragraph(
        "Five algorithms were trained using 5-fold Stratified Cross-Validation and hyperparameter tuning via GridSearchCV. "
        "The models were evaluated on an unseen 20% test set."
    )
    
    # Simple table representation
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Model'
    hdr_cells[1].text = 'Accuracy'
    hdr_cells[2].text = 'Precision'
    hdr_cells[3].text = 'Recall'
    hdr_cells[4].text = 'F1-Score'
    
    data = [
        ('Logistic Regression', '97.50%', '100.0%', '96.00%', '97.96%'),
        ('KNN', '97.50%', '100.0%', '96.00%', '97.96%'),
        ('Decision Tree', '91.25%', '93.88%', '92.00%', '92.93%'),
        ('Random Forest', '98.75%', '100.0%', '98.00%', '98.99%'),
        ('Gradient Boosting', '96.25%', '97.96%', '96.00%', '96.97%')
    ]
    for row in data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val

    doc.add_paragraph("\nComparative Visualizations:")
    eval_images = ["09_confusion_matrices.png", "10_roc_curves.png"]
    for img in eval_images:
        img_path = os.path.join(PLOTS_DIR, img)
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(5.5))

    doc.add_heading('6. Final Analysis & Answers', level=1)
    doc.add_paragraph(
        "Q1: Can kidney-disease-related classification be performed using ML?\n"
        "Yes. The models demonstrated high discriminative power, all scoring above 91% accuracy.\n\n"
        "Q2: Which features are most important?\n"
        "Hemoglobin, Packed Cell Volume, Specific Gravity, and Serum Creatinine.\n\n"
        "Q3: Which algorithm performs best?\n"
        "Random Forest outperformed others with 98.75% accuracy and an F1-score of 98.99%.\n\n"
        "Q4: How does missing-value treatment affect the model?\n"
        "Proper imputation (KNN & Mode) allowed retention of 100% of the data without bias, enabling robust model training.\n\n"
        "Q5: Which model has the best recall?\n"
        "Random Forest achieved the best recall (98.00%), which is crucial in healthcare to minimize false negatives (missed diagnoses).\n\n"
        "Q6: Can the model be deployed successfully?\n"
        "Yes, the Random Forest model was successfully serialized and deployed via a Streamlit web application."
    )

    doc.save(DOCX_PATH)
    print(f"✅ Generated {DOCX_PATH}")


# ==========================================
# Generate PPTX Presentation
# ==========================================
def create_pptx_presentation():
    prs = Presentation()
    
    # Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Chronic Kidney Disease Classification"
    subtitle.text = "Machine Learning Case Study\nUsing Clinical Measurements"
    
    # Slide 1: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Problem Statement & Objectives"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Goal: Classify patients as having CKD or Not CKD using ML."
    p = tf.add_paragraph()
    p.text = "Objectives:"
    p.level = 1
    bullet_points = [
        "Clean and preprocess clinical data",
        "Handle significant missing values",
        "Identify key disease predictors",
        "Compare 5 ML algorithms",
        "Deploy the best model via Web App"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 2

    # Slide 2: Data & Preprocessing
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Dataset & Preprocessing"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "UCI CKD Dataset: 400 Patients, 24 Features"
    p = tf.add_paragraph()
    p.text = "Challenge: Up to 38% missing data in some features."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Solution: Robust Imputation Strategy"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Numerical: KNN Imputation (k=5)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Categorical: Mode Imputation"
    p.level = 2

    # Slide 3: Feature Importance
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title only
    title = slide.shapes.title
    title.text = "Key Clinical Predictors"
    img_path = os.path.join(PLOTS_DIR, "08_feature_importance.png")
    if os.path.exists(img_path):
        prs.slides[-1].shapes.add_picture(img_path, PptxInches(1), PptxInches(1.5), width=PptxInches(8))

    # Slide 4: Model Performance
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Model Evaluation & Comparison"
    img_path = os.path.join(PLOTS_DIR, "11_metrics_comparison.png")
    if os.path.exists(img_path):
        prs.slides[-1].shapes.add_picture(img_path, PptxInches(1.5), PptxInches(1.5), width=PptxInches(7))

    # Slide 5: Confusion Matrices
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Confusion Matrices (False Positives/Negatives)"
    img_path = os.path.join(PLOTS_DIR, "09_confusion_matrices.png")
    if os.path.exists(img_path):
        prs.slides[-1].shapes.add_picture(img_path, PptxInches(0.2), PptxInches(2), width=PptxInches(9.5))

    # Slide 6: Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Conclusions"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Random Forest is the winning model."
    
    p = tf.add_paragraph()
    p.text = "Performance: 98.75% Accuracy, 98% Recall"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Why Recall matters: Minimizes False Negatives (missing a sick patient)."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Deployment: Model successfully integrated into a Streamlit Web Application for real-time clinical use."
    p.level = 1

    prs.save(PPTX_PATH)
    print(f"✅ Generated {PPTX_PATH}")

if __name__ == "__main__":
    create_docx_report()
    create_pptx_presentation()
